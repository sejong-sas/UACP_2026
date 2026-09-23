from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path("/tmp/uacp_extract/template.docx")
OUT = ROOT / "paper_draft_v1.docx"

def font(run, size=9.5, bold=False, italic=False):
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic

def clear_body(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)

def add_para(doc, text="", style="Body", align=None, before=0, after=3, indent=True, size=9.5):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.12
    if indent and style == "Body":
        p.paragraph_format.first_line_indent = Pt(10)
    if text:
        r = p.add_run(text)
        font(r, size)
    return p

def heading(doc, text, level=1):
    p = doc.add_paragraph(style="HeadingCustom")
    p.paragraph_format.space_before = Pt(9 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    font(r, 11.5 if level == 1 else 10.5, True)
    return p

def cell_text(cell, value, bold=False, size=7.6):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(value))
    font(r, size, bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def shade(cell, fill="D9EAF7"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def table(doc, headers, rows, size=7.6):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if "Table Grid" in [s.name for s in doc.styles]:
        t.style = "Table Grid"
    else:
        tblPr = t._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            tag = OxmlElement("w:" + edge)
            tag.set(qn("w:val"), "single")
            tag.set(qn("w:sz"), "4")
            tag.set(qn("w:space"), "0")
            tag.set(qn("w:color"), "B7B7B7")
            borders.append(tag)
        tblPr.append(borders)
    for i, x in enumerate(headers):
        cell_text(t.rows[0].cells[i], x, True, size)
        shade(t.rows[0].cells[i])
    for row in rows:
        cells = t.add_row().cells
        for i, x in enumerate(row):
            cell_text(cells[i], x, False, size)
    col_width = Inches(3.0 / len(headers))
    for column in t.columns:
        column.width = col_width
    grid = t._tbl.tblGrid
    for grid_col in grid.gridCol_lst:
        grid_col.set(qn("w:w"), str(int(3.0 * 1440 / len(headers))))
    for row in t.rows:
        for cell in row.cells:
            cell.width = col_width
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

def caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    font(r, 8.5)
    return p

def image(doc, path, width=3.15):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))

doc = Document(str(TEMPLATE))
clear_body(doc)
for s in doc.sections:
    s.top_margin = Inches(0.55)
    s.bottom_margin = Inches(0.45)
    s.left_margin = Inches(0.65)
    s.right_margin = Inches(0.65)

for name in ["Body", "HeadingCustom", "RefCustom"]:
    if name not in doc.styles:
        doc.styles.add_style(name, 1)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Evidential MIMO 채널 예측에서\nParameter-Efficient Adaptation의 성능과 비용 Trade-off 분석")
font(r, 15, True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Performance–Cost Trade-off of Parameter-Efficient Adaptation for Evidential MIMO Channel Prediction")
font(r, 10.5, True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("[저자명 확인 필요]¹")
font(r, 10, True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("¹[소속 및 이메일 확인 필요]")
font(r, 8.8)

heading(doc, "요       약")
add_para(doc, "Uncertainty-Aware Channel Prediction in MIMO Networks (UACP)는 sparse Channel Frequency Response (CFR)로부터 전체 CFR을 복원하고, Epistemic Uncertainty가 높아질 때 학습 분포 밖의 채널로 판단하여 adaptation을 유도한다. 그러나 실제 adaptation에서 전체 parameter를 업데이트할지 일부 parameter만 업데이트할지에 대한 비용과 성능의 비교는 구체적으로 제시되지 않았다. 본 연구는 UACP를 새 알고리즘으로 제안하지 않고, evidential MIMO channel predictor에서 parameter-efficient adaptation이 가능한 범위와 한계를 실험적으로 분석한다. Epoch 3 pretrained model을 동일한 120 ns adaptation protocol에서 3.16% Partial Fine-Tuning, 25.03% Partial Fine-Tuning, Full Fine-Tuning으로 업데이트하고, 120 ns reconstruction 회복, 120 ns Epistemic Uncertainty 변화, 20 ns/80 ns ID 성능 보존, adaptation에 사용하지 않은 1 ms Far-OOD uncertainty, trainable parameter 수·training time·GPU memory를 비교하였다. 3.16%는 Full 120 ns NMSE 개선량의 약 87.7–88.9%를 얻으면서 parameter 수와 peak VRAM을 크게 줄였다. 25%는 120 ns reconstruction과 120 ns Epistemic 감소가 Full과 유사한 수준이었지만 1 ms Epistemic Uncertainty가 크게 낮아졌다. Full은 120 ns Epistemic 감소가 가장 컸으나 ID forgetting과 비용도 더 컸다. 따라서 현재 구현과 고정된 adaptation schedule에서는 reconstruction 성능과 uncertainty 보존 및 비용이 같은 방향으로 개선되지 않는 trade-off가 확인되었다.")
add_para(doc, "핵심어: MIMO, Channel State Information, Evidential Regression, Out-of-Distribution, Partial Fine-Tuning, Epistemic Uncertainty", indent=False, after=8)

heading(doc, "1. 서론")
for text in [
"Multiple-Input Multiple-Output (MIMO) 시스템은 여러 안테나를 이용해 spatial multiplexing과 beamforming 이득을 얻지만, 송신기에서 precoding을 수행하려면 Channel State Information (CSI)을 확보해야 한다. 특히 Frequency Division Duplex (FDD) 환경에서는 수신기가 추정한 downlink CSI를 feedback link를 통해 송신기로 전달해야 하므로, 안테나 수와 subcarrier 수가 증가할수록 feedback overhead가 커진다 [1]. 이 문제를 줄이기 위해 CSI를 codeword로 압축하거나 일부 subcarrier만 전송한 뒤 나머지를 복원하는 방법이 연구되어 왔다 [2, 3].",
"Sparse feedback을 이용한 복원에서는 채널의 주파수 선택성이 중요한 변수이다. 채널이 주파수 방향으로 완만하게 변하면 적은 수의 관측 subcarrier만으로도 나머지를 추정할 수 있지만, delay spread가 커져 coherence bandwidth가 작아지면 인접 subcarrier 사이의 변화가 커져 복원이 어려워진다. 따라서 고정된 sampling rate만 사용하는 방식은 채널 환경 변화에 따라 복원 품질과 feedback 양 사이의 균형을 조절하기 어렵다.",
"UACP는 sparse CFR을 입력으로 받아 전체 CFR과 함께 prediction uncertainty를 출력하는 evidential channel predictor를 사용한다 [4]. UACP는 Normal-Inverse-Wishart (NIW) evidential distribution에서 Aleatoric Uncertainty와 Epistemic Uncertainty를 분리한다. Aleatoric Uncertainty는 채널 자체의 예측 난이도와 관련되고, Epistemic Uncertainty는 모델이 충분히 학습하지 못한 환경과 관련된 신호로 해석된다. 원 논문은 Epistemic Uncertainty가 높아지면 full feedback을 요청하고 model adaptation을 trigger하는 흐름을 제시한다 [4].",
"하지만 adaptation을 trigger한 뒤 실제 predictor의 parameter를 어떤 범위로 업데이트해야 하는지는 별도의 문제이다. 전체 parameter를 업데이트하면 target channel에 빠르게 맞출 수 있지만 계산량이 커지고 기존 환경의 성능 또는 uncertainty boundary가 변할 수 있다. 반대로 일부 parameter만 업데이트하면 비용을 줄일 수 있지만 새로운 channel의 reconstruction과 evidential output을 충분히 조정하지 못할 수 있다. Adapter tuning과 LoRA 같은 parameter-efficient transfer 연구는 적은 수의 parameter를 학습해 full fine-tuning과 다른 비용 구조를 만드는 방법을 보여주었지만 [5, 6], evidential MIMO channel prediction에서 reconstruction·uncertainty·Far-OOD 보존을 함께 비교한 것은 별도의 검토가 필요하다.",
"이에 본 연구는 다음 질문을 둔다. Evidential channel predictor가 120 ns OOD channel에서 adaptation을 수행할 때 전체 parameter를 업데이트하지 않고 3.16% 또는 25%만 업데이트해도 reconstruction을 회복할 수 있으며, 그 비용 절감이 ID 성능과 Far-OOD Epistemic Uncertainty 보존으로 이어지는가? 본 연구의 목적은 Partial Fine-Tuning 자체를 새로운 알고리즘으로 제안하는 것이 아니라 고정된 구현과 protocol에서 adaptation 성능과 비용의 trade-off를 확인하는 것이다."
]:
    add_para(doc, text)

heading(doc, "2. 실험 디자인")
heading(doc, "2.1 데이터셋 및 채널 환경", 2)
add_para(doc, "원 UACP 논문은 Sionna 기반 MIMO OFDM simulation에서 carrier frequency 3.5 GHz, subcarrier spacing 30 kHz, SNR 15 dB, 100,000개의 training CFR, [10, 100] ns uniform delay spread, learning rate 10⁻⁴, batch size 4096, 150 epochs를 사용한다고 기술한다 [4]. 현재 repository에서는 이 설정 중 일부를 유지했지만 hardware와 구현 제약 때문에 학습 규모와 observation pipeline이 완전히 같지 않다.")
table(doc, ["항목","원 논문 명시","현재 구현","구분"], [
["MIMO / K","MIMO OFDM / 식에서 정의","2×2 MIMO, K=1024","코드·config"],
["채널","Sionna 사용","Sionna 2.0.1, 3GPP TDL-A","TDL-A 가정"],
["주파수","3.5 GHz, 30 kHz","3.5 GHz, 30 kHz","공통"],
["mobility / normalization","세부 확인 필요","zero mobility, normalize=false","구현 가정"],
["training","U[10,100] ns, 100k","Epoch 3 checkpoint는 100k×5","현재 baseline"],
["observation","sparse CFR + mask","9-channel, Ng 4/8/16/32","세부 구현"],
["noise","15 dB SNR","reported CFR에만 AWGN, clean target","단계 가정"],
["Ψ","full covariance L Lᵀ","diagonal Ψ approximation","근사"],
], size=7.1)
add_para(doc, "본 연구의 ID-Easy는 20 ns, ID-Hard는 80 ns이다. 120 ns는 adaptation target이면서 Near-OOD 평가 환경으로 사용했고, 1 ms는 adaptation 학습에 사용하지 않은 Far-OOD test 환경으로 유지했다. 120 ns adaptation data는 train 1,000개와 validation 200개로 구성했으며 각 test regime은 1,000개를 사용했다. 이 split 수와 3-epoch budget은 원 논문 설정이 아니라 본 연구의 연구 확장이다.")

heading(doc, "2.2 Evidential Channel Predictor", 2)
for text in [
"입력은 2×2 MIMO의 네 antenna-pair CFR을 real/imaginary channel 8개로 펼친 뒤 관측되지 않은 subcarrier를 0으로 채우고 binary mask 1개를 추가한 [B, 9, 1024] tensor이다. predictor는 192 hidden channels의 1D convolution backbone과 32개 residual block을 사용한다. 출력은 CFR prediction mean γ와 evidential parameter κ, Ψ, ν이다. 원 UACP의 antenna-pair factorization과 NIW formulation은 [4]를 따른다.",
"현재 구현은 pair-level scalar κ와 ν, subcarrier별 diagonal Ψ를 사용한다. 따라서 다음 식은 원 논문의 full Ψ 식을 그대로 구현한 것이 아니라 current diagonal specialization의 계산식이다 [4]."
]:
    add_para(doc, text)
add_para(doc, "Σᵃˡᵉ = Ψ / (ν − 2K − 1),    Σᵉᵖⁱ = Σᵃˡᵉ / κ", align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, after=2)
add_para(doc, "평가에서는 각 antenna pair covariance의 real/imaginary trace를 계산하고 네 pair 평균을 낸 뒤 omitted subcarrier 평균을 사용했다. Evidence regularizer가 포함된 loss는 L = L_NLL + λ_reg L_reg이고 현재 adaptation에서는 λ_reg=10⁻³을 사용했다 [4]. Evidential regression은 한 번의 forward pass에서 예측값과 uncertainty parameter를 함께 출력할 수 있다는 장점이 있으나 uncertainty 값이 calibration과 parameter coupling에 영향을 받는다 [7, 8].")

heading(doc, "2.3 Partial Fine-Tuning", 2)
add_para(doc, "모든 비교는 Epoch 3 pretrained checkpoint에서 시작했다. adaptation 학습은 120 ns full CFR로 clean target을 만들되 predictor 입력은 15 dB AWGN이 적용된 sparse CFR과 mask로 구성했다. 3.16%, 25%, Full은 동일한 sample order, mask/noise schedule, Adam, learning rate 10⁻⁴, batch size 8, 3 epochs, 375 optimizer steps를 사용했다. 120 ns test set은 adaptation train/validation에 재사용하지 않았고 cross-split CFR byte-hash duplicate는 0이었다.")
table(doc, ["모델","업데이트 layer","Trainable params","비율"], [
["Pre","없음","0","0%"],
["Partial 3.16%","ResidualBlock 31 + four evidential heads","373,656","3.1625%"],
["Partial 25.03%","ResidualBlock 24–31 + four heads","2,956,824","25.0253%"],
["Full","전체 predictor","11,815,320","100%"],
], size=7.7)
image(doc, ROOT / "output/UACP_Partial_FineTuning_Experiment_Report_20260923_v5/figures/fig2_adaptation_flow.png")
caption(doc, "그림 1. 120 ns full feedback은 clean target을 만드는 데 사용되지만 adaptation 입력은 sparse CFR과 mask로 유지한 현재 protocol.")
add_para(doc, "이 layer policy는 parameter-efficient transfer learning의 문제의식과 관련되지만 [5, 6], 본 연구가 adapter module 또는 low-rank update를 새로 제안하는 것은 아니다. 기존 residual block과 evidential head 중 일부의 update scope를 비교한다.")

heading(doc, "2.4 평가 방법", 2)
for x in [
"120 ns reconstruction recovery: omitted-subcarrier NMSE(dB)를 측정하고 improvement = NMSE_Pre − NMSE_Post로 정의했다. 양수일수록 adaptation 후 NMSE가 낮아진다.",
"Epistemic adaptation: 120 ns Epistemic의 post-minus-pre 변화를 사용했다. 음수일수록 target 환경에서 model uncertainty가 감소한 것이다.",
"ID forgetting: 20 ns와 80 ns에서 ΔNMSE = NMSE_Post − NMSE_Pre를 사용했다. 양수는 adaptation 후 성능 저하를 뜻한다.",
"Far-OOD preservation: 학습하지 않은 1 ms에서 Epistemic 절대값과 120 ns 대비 ratio를 비교했다. AUROC만으로는 uncertainty magnitude collapse를 놓칠 수 있어 보조 지표로만 해석했다.",
"Adaptation cost: trainable parameter 수, wall-clock training time, seconds/step, peak allocated VRAM을 기록했다."
]:
    p = add_para(doc, "• " + x, indent=False, after=2, size=9.2)
    p.paragraph_format.left_indent = Pt(12)

heading(doc, "3. 결과 분석")
heading(doc, "3.1 Adaptation 전 모델의 동작", 2)
add_para(doc, "본 연구의 목적은 UACP 전체 reproduction이 아니라 adaptation scope 비교이므로 baseline 재현은 최소 조건만 확인한다. Epoch 3 checkpoint의 공통 evaluation에서 Ng=16의 omitted NMSE는 20 ns −20.9567 dB, 80 ns −18.7954 dB, 120 ns −16.8910 dB, 1 ms 1.6995 dB였다. Ng=32에서도 20 ns −18.1301 dB, 80 ns −15.9997 dB, 120 ns −14.2149 dB, 1 ms 1.8416 dB로 80 ns가 20 ns보다 어려웠고 1 ms에서 reconstruction이 크게 악화되었다.")
add_para(doc, "Epistemic Uncertainty도 120 ns에서 20/80 ns보다 높고 1 ms에서 크게 증가했다. 예를 들어 Ng=16에서 120 ns는 0.008393, 1 ms는 32985.8이었다. 따라서 이 checkpoint는 CFR reconstruction, ID 내부 난이도 차이, OOD Epistemic 증가, adaptation 수행이라는 최소 조건을 만족한다. 다만 ID-Hard의 Aleatoric 증가와 Near-OOD Epistemic separation이 모든 baseline 실험에서 원 UACP 논문처럼 일관되게 나타난 것은 아니므로 UACP를 완전히 재현했다고 표현하지 않는다.")

heading(doc, "3.2 OOD 채널 적응 성능", 2)
add_para(doc, "동일한 120 ns adaptation protocol에서 얻은 epoch-3 omitted NMSE와 Pre 대비 improvement를 비교하였다. Ng=16에서 3.16%는 1.2428 dB, 25%는 1.5210 dB, Full은 1.4167 dB 개선되었다. Ng=32에서는 각각 0.9180 dB, 1.1287 dB, 1.0329 dB였다.")
table(doc, ["Scope","120 ns NMSE Ng16","개선 Ng16","120 ns NMSE Ng32","개선 Ng32","Full 대비"], [
["Pre","-16.8910","—","-14.2149","—","—"],
["Partial 3.16%","-18.1338","1.2428","-15.1330","0.9180","87.7% / 88.9%"],
["Partial 25.03%","-18.4120","1.5210","-15.3437","1.1287","107.4% / 109.3%"],
["Full","-18.3077","1.4167","-15.2478","1.0329","100% / 100%"],
], size=7.4)
add_para(doc, "3.16%는 전체 parameter의 약 3%만 업데이트하면서도 Full improvement의 87.7–88.9%에 도달했다. 이는 작은 scope에서도 target channel의 reconstruction 함수를 조정할 수 있음을 보여준다. 다만 25%가 3.16%보다 수치상 큰 improvement를 보였다는 사실만으로 25%가 더 좋은 adaptation policy라고 결론내릴 수는 없다. 세 scope의 동일한 3-epoch 조건에 한정된 관찰이며 uncertainty와 Far-OOD 결과를 함께 보아야 한다.")

heading(doc, "3.3 Uncertainty 변화 및 기존 성능 보존", 2)
table(doc, ["Scope","120 ns Epi Ng16","1 ms Epi Ng16","120 ns Epi Ng32","1 ms Epi Ng32"], [
["Pre","0.008393","32985.8","0.053175","3.0569e6"],
["Partial 3.16%","0.008263","27697.6","0.046497","2.2521e6"],
["Partial 25.03%","0.007924","22.10","0.026943","82.30"],
["Full","0.007759","353.28","0.026344","1195.07"],
], size=7.4)
for text in [
"120 ns Epistemic 변화는 Ng=16에서 3.16% −0.000130, 25% −0.000469, Full −0.000633이고 Ng=32에서는 −0.006678, −0.026232, −0.026831이었다. 따라서 Full이 target environment에서 Epistemic을 가장 크게 낮추었다.",
"25%는 120 ns에서는 Full과 가까운 감소를 보였지만 1 ms에서는 Epistemic magnitude가 크게 줄었다. Ng=16 기준 Pre의 1 ms Epistemic 32985.8이 25%에서 22.10으로, Full에서 353.28로 감소했다. Ng=32에서도 3.0569e6이 25%에서 82.30, Full에서 1195.07로 낮아졌다. 3.16%는 1 ms uncertainty가 27697.6과 2.2521e6으로 Pre에 상대적으로 가까웠다.",
"이 결과는 reconstruction recovery와 uncertainty preservation이 동일한 방향으로 움직이지 않음을 보여준다. 25%는 target-domain accuracy만 보면 좋은 결과이지만 학습하지 않은 1 ms에 대해서는 model이 과도하게 확신하는 방향으로 Epistemic scale이 이동했다. 따라서 Partial Fine-Tuning을 평가할 때 target-domain accuracy 하나만 사용하는 것은 충분하지 않다.",
"ID forgetting은 20 ns에서 Ng=16 기준 3.16% +1.3651 dB, 25% +2.4718 dB, Full +2.7708 dB였고 Ng=32에서는 +1.3094, +2.2528, +2.5139 dB였다. 80 ns에서는 Ng=16 기준 −0.0758, +0.0275, +0.1950 dB, Ng=32 기준 −0.1127, +0.0689, +0.2261 dB였다. 즉 3.16%는 Full보다 ID-Easy forgetting이 작았고 80 ns에서는 NMSE가 소폭 개선되었다. 25%는 120 ns reconstruction을 더 크게 개선했지만 3.16%보다 20 ns forgetting이 커졌다. Sequential adaptation에서 이전 능력이 손상될 수 있는 문제는 catastrophic forgetting 연구에서도 다루어졌다 [9]."
]:
    add_para(doc, text)

heading(doc, "3.4 적응 효율성", 2)
table(doc, ["Scope","Trainable params","비율","Time(s)","sec/step","Peak VRAM(MiB)"], [
["Partial 3.16%","373,656","3.1625%","137.32","0.3035","117.84"],
["Partial 25.03%","2,956,824","25.0253%","149.34","0.3431","275.55"],
["Full","11,815,320","100%","285.94","0.6982","816.16"],
], size=7.4)
add_para(doc, "3.16%는 Full보다 trainable parameter가 96.8% 적고 wall-clock time은 48.0%, peak allocated VRAM은 14.4%였다. 25%는 parameter 비율이 25.0%였지만 time은 Full의 52.2%, VRAM은 33.8%였다. 따라서 trainable parameter 수와 실제 wall-clock/VRAM은 같은 비율로 감소하지 않았다. Optimizer state뿐 아니라 frozen backbone의 forward와 activation 저장이 남아 있기 때문이다.")
add_para(doc, "성능과 비용을 같이 보면 3.16%는 가장 작은 cost로 상당한 reconstruction recovery와 상대적으로 큰 Far-OOD uncertainty 보존을 보였다. 25%는 reconstruction과 120 ns uncertainty adaptation을 더 크게 수행했지만 Far-OOD uncertainty 보존에서 불리했다. Full은 120 ns uncertainty를 가장 많이 줄였으나 가장 큰 ID forgetting과 cost를 보였다. 따라서 이 실험에서는 하나의 scope가 모든 지표에서 우세하지 않았고 adaptation 목적에 따라 선택이 달라질 수 있다.")

heading(doc, "4. 결론")
for text in [
"본 연구는 UACP 재현 자체가 아니라 evidential MIMO channel predictor가 120 ns OOD 환경에서 adaptation할 때 Partial Fine-Tuning과 Full Fine-Tuning 사이에 어떤 trade-off가 생기는지 분석했다. 동일한 Epoch 3 checkpoint와 120 ns adaptation schedule에서 3.16% Partial Fine-Tuning은 Full 120 ns NMSE 개선량의 약 87.7–88.9%를 얻었고 parameter 수·training time·peak VRAM을 줄였다. 25% Partial Fine-Tuning은 3.16%보다 120 ns reconstruction 개선과 Epistemic 감소가 컸으며 Full과 유사한 수준에 도달했다. 그러나 25%와 Full은 1 ms Far-OOD Epistemic magnitude가 크게 감소했고 Full은 ID forgetting과 cost도 가장 컸다.",
"따라서 현재 결과는 Partial Fine-Tuning으로 어느 정도의 OOD reconstruction adaptation이 가능하다는 점과 parameter 범위·adaptation 성능·기존 성능·Far-OOD uncertainty·계산 비용 사이에 단순하지 않은 trade-off가 있다는 점을 보여준다. 특히 25% 결과는 reconstruction만 보고 성공으로 판단할 수 없음을 보여준다. 본 연구가 증명한 범위는 현재의 diagonal Ψ approximation, sparse observation과 noise pipeline, dataset, optimizer, learning rate, adaptation sample 수, epoch 수, layer scope에 한정된다.",
"한계로는 첫째 원 논문과 현재 구현의 full covariance 및 noise stage가 동일하지 않다. 둘째 주 비교는 3 epochs와 단일 fixed protocol에 기반하며 multi-seed와 longer online adaptation 검증은 제한적이다. 셋째 repository Fig.11 controller는 다른 checkpoint와 delay-sweep provenance를 사용하므로 본 adaptation 비교의 직접적인 runtime 결과로 사용하지 않았다. 넷째 ID 내부 Aleatoric Uncertainty 증가 경향이 모든 실험에서 원 연구만큼 명확하게 재현되지는 않았다. 향후에는 uncertainty response나 feature drift로 update layer를 선택하는 방법, uncertainty head를 고정하거나 별도 regularization하는 방법, multi-seed·multiple OOD delay·online adaptation을 추가로 비교할 필요가 있다."
]:
    add_para(doc, text)

heading(doc, "참고문헌")
refs = [
"[1] A. Goldsmith, Wireless Communications, Cambridge University Press, 2005.",
"[2] C.-K. Wen, W.-T. Shih, and S. Jin, “Deep Learning for Massive MIMO CSI Feedback,” IEEE Wireless Communications Letters, vol. 7, no. 5, pp. 748–751, 2018. doi:10.1109/LWC.2018.2818160.",
"[3] J. Wang, G. Gui, T. Ohtsuki, B. Adebisi, H. Gacanin, and H. Sari, “Compressive Sampled CSI Feedback Method Based on Deep Learning for FDD Massive MIMO Systems,” IEEE Transactions on Communications, vol. 69, no. 9, pp. 5873–5885, 2021. doi:10.1109/TCOMM.2021.3086525.",
"[4] UACP authors, “UACP: Uncertainty-Aware Channel Prediction in MIMO Networks,” submitted to ACM MobiHoc ’26, 2026. Repository PDF: mobihoc26-paper289.pdf. DOI 확인 필요.",
"[5] N. Houlsby et al., “Parameter-Efficient Transfer Learning for NLP,” Proceedings of the 36th International Conference on Machine Learning, PMLR 97, pp. 2790–2799, 2019.",
"[6] E. J. Hu et al., “LoRA: Low-Rank Adaptation of Large Language Models,” International Conference on Learning Representations, 2022.",
"[7] A. Amini, W. Schwarting, A. Soleimany, and D. Rus, “Deep Evidential Regression,” Advances in Neural Information Processing Systems 33, pp. 14927–14937, 2020.",
"[8] N. Meinert and A. Lavin, “Multivariate Deep Evidential Regression,” arXiv:2104.06135, 2021.",
"[9] J. Kirkpatrick et al., “Overcoming Catastrophic Forgetting in Neural Networks,” Proceedings of the National Academy of Sciences, vol. 114, no. 13, pp. 3521–3526, 2017. doi:10.1073/pnas.1611835114.",
]
for ref in refs:
    p = add_para(doc, ref, style="RefCustom", indent=False, after=2, size=8.5)
    p.paragraph_format.left_indent = Pt(8)
    p.paragraph_format.first_line_indent = Pt(-8)

doc.core_properties.title = "Evidential MIMO 채널 예측에서 Parameter-Efficient Adaptation의 성능과 비용 Trade-off 분석"
doc.core_properties.subject = "KIPS 학술발표대회 논문 초안"
doc.core_properties.author = "[저자명 확인 필요]"
doc.save(str(OUT))
print(OUT)
